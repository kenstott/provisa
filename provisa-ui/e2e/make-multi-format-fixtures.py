#!/usr/bin/env python3
# Copyright (c) 2026 Kenneth Stott
# Canary: 7b6f0a3d-2e4c-4a1b-9d5e-8c6f1a2b3d4e
# Canary: placeholder
#
# This source code is licensed under the Business Source License 1.1
# found in the LICENSE file in the root directory of this source tree.
#
# NOTICE: Use of this software for training artificial intelligence or
# machine learning models is strictly prohibited without explicit written
# permission from the copyright holder.

"""REQ-1690: writes one small table-bearing file per format Calcite's file adapter recognizes
(FileSchema.java's CONVERTIBLE_EXTENSIONS/TABLE_SOURCE_EXTENSIONS: csv, json, parquet native;
xlsx, md, docx, pptx, html converted to JSON first), plus a subdirectory file (recursive
discovery) and a .jsonl/.xml pair that the SAME adapter does NOT recognize — confirmed absent
from both extension sets — so file-connector-multi-format.spec.ts can assert they are correctly
excluded rather than silently mis-scanned.

Usage: python make-multi-format-fixtures.py <output-dir>
Writes the fixtures directly into <output-dir> (plus <output-dir>/subdir/) and prints nothing;
the caller already knows the directory it passed in.
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path


def _write_csv(path: Path) -> None:
    path.write_text(
        "id,name,email\n"
        "1,Alice,alice@example.com\n"
        "2,Bob,bob@example.com\n"
        "3,Carol,carol@example.com\n"
    )


def _write_json(path: Path) -> None:
    path.write_text(
        json.dumps(
            [
                {"id": 1, "total": 100.5, "status": "shipped"},
                {"id": 2, "total": 42.0, "status": "pending"},
                {"id": 3, "total": 7.25, "status": "shipped"},
            ]
        )
    )


def _write_jsonl(path: Path) -> None:
    """NOT a recognized Calcite file-adapter extension (confirmed absent from FileSchema.java's
    extension sets) — written so the test can assert it is correctly excluded, not just missing."""
    lines = [
        {"id": 1, "event": "login"},
        {"id": 2, "event": "logout"},
    ]
    path.write_text("\n".join(json.dumps(line) for line in lines) + "\n")


def _write_xml(path: Path) -> None:
    """Same as jsonl: not in FileSchema.java's extension sets — written to prove exclusion."""
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<config>\n"
        '  <setting name="retries" value="3"/>\n'
        '  <setting name="timeout" value="30"/>\n'
        "</config>\n"
    )


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None  # a freshly constructed Workbook always has an active sheet
    ws.title = "Sheet1"
    ws.append(["sku", "quantity", "price"])
    ws.append(["WIDGET-1", 10, 2.5])
    ws.append(["WIDGET-2", 5, 9.99])
    wb.save(path)


def _write_parquet(path: Path) -> None:
    import pyarrow as pa
    import pyarrow.parquet as pq

    table = pa.table(
        {
            "region": pa.array(["east", "west", "north"]),
            "revenue": pa.array([1000.0, 2500.5, 800.25]),
        }
    )
    pq.write_table(table, path)


def _write_md(path: Path) -> None:
    path.write_text("# Notes\n\n| id | label |\n|----|-------|\n| 1  | first |\n| 2  | second |\n")


def _write_html(path: Path) -> None:
    path.write_text(
        "<html><body>\n"
        "<table>\n"
        "<tr><th>id</th><th>city</th></tr>\n"
        "<tr><td>1</td><td>Austin</td></tr>\n"
        "<tr><td>2</td><td>Denver</td></tr>\n"
        "</table>\n"
        "</body></html>\n"
    )


def _write_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rows, cols = 3, 2
    table = slide.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "id"
    table.cell(0, 1).text = "team"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "Alpha"
    table.cell(2, 0).text = "2"
    table.cell(2, 1).text = "Beta"
    prs.save(str(path))


_DOCX_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)

_DOCX_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/>'
    "</Relationships>"
)

_DOCX_DOCUMENT = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:body>"
    "<w:tbl>"
    "<w:tr><w:tc><w:p><w:r><w:t>id</w:t></w:r></w:p></w:tc>"
    "<w:tc><w:p><w:r><w:t>name</w:t></w:r></w:p></w:tc></w:tr>"
    "<w:tr><w:tc><w:p><w:r><w:t>1</w:t></w:r></w:p></w:tc>"
    "<w:tc><w:p><w:r><w:t>Alpha</w:t></w:r></w:p></w:tc></w:tr>"
    "<w:tr><w:tc><w:p><w:r><w:t>2</w:t></w:r></w:p></w:tc>"
    "<w:tc><w:p><w:r><w:t>Beta</w:t></w:r></w:p></w:tc></w:tr>"
    "</w:tbl>"
    "</w:body>"
    "</w:document>"
)


def _write_docx(path: Path) -> None:
    """A minimal hand-built OOXML package (no python-docx dependency): one table, which is all
    Calcite's POI-backed docx conversion needs to produce a table."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _DOCX_CONTENT_TYPES)
        z.writestr("_rels/.rels", _DOCX_RELS)
        z.writestr("word/document.xml", _DOCX_DOCUMENT)


def main() -> None:
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    subdir = out / "subdir"
    subdir.mkdir(parents=True, exist_ok=True)

    _write_csv(out / "customers.csv")
    _write_json(out / "orders.json")
    _write_jsonl(out / "events.jsonl")
    _write_xml(out / "config.xml")
    _write_xlsx(out / "products.xlsx")
    _write_parquet(out / "sales.parquet")
    _write_md(out / "notes.md")
    _write_html(out / "page.html")
    _write_pptx(out / "deck.pptx")
    _write_docx(out / "memo.docx")
    _write_csv(subdir / "extra.csv")


if __name__ == "__main__":
    main()
