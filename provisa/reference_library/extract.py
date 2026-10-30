# Copyright (c) 2026 Kenneth Stott
# Canary: 499b6572-057d-4b02-bbdf-158725a6e032
#
# This source code is licensed under the Business Source License 1.1
# found in the LICENSE file in the root directory of this source tree.
#
# NOTICE: Use of this software for training artificial intelligence or
# machine learning models is strictly prohibited without explicit written
# permission from the copyright holder.

"""A library document read into sections (REQ-2185).

A section is a run of paragraphs under one heading, or on one page, so a passage cut from it
cites where it sits. HTML is read by its headings (``h1``-``h6``), markdown by its ``#`` headings,
a Word document by its heading styles (python-docx), a PDF by its pages (pypdf), plain text as
paragraphs. A media type this module does not read is refused by name.
"""

# Requirements: REQ-2185

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import PurePosixPath

HTML = "text/html"
MARKDOWN = "text/markdown"
PLAIN = "text/plain"
PDF = "application/pdf"
WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

#: The media types a library document may have, by the file extension that names each.
EXTENSIONS: dict[str, str] = {
    ".html": HTML,
    ".htm": HTML,
    ".md": MARKDOWN,
    ".markdown": MARKDOWN,
    ".txt": PLAIN,
    ".pdf": PDF,
    ".docx": WORD,
}

READABLE: frozenset[str] = frozenset(EXTENSIONS.values())


class UnreadableDocument(ValueError):
    """A document the library cannot read, said naming why."""


@dataclass
class Section:
    """A run of paragraphs under one heading (``heading``) or on one page (``page``)."""

    heading: str | None
    page: int | None
    paragraphs: list[str] = field(default_factory=list)


@dataclass
class Document:
    """A document as read: its title (None when it states none) and its sections."""

    title: str | None
    sections: list[Section]


def media_type_of_name(name: str) -> str:
    """The media type a file's name says it has, or :class:`UnreadableDocument`."""
    suffix = PurePosixPath(name).suffix.lower()
    media = EXTENSIONS.get(suffix)
    if media is None:
        raise UnreadableDocument(
            f"{name!r} is not a document the library reads; it reads "
            f"{', '.join(sorted(EXTENSIONS))} files"
        )
    return media


def media_type_of_response(content_type: str | None, url: str) -> tuple[str, str | None]:
    """The media type and charset an HTTP response declares (``Content-Type``), or, when it
    declares none, the one its URL's path names."""
    if content_type:
        parts = [p.strip() for p in content_type.split(";")]
        media = parts[0].lower()
        charset = next(
            (
                p.split("=", 1)[1].strip().strip('"')
                for p in parts[1:]
                if p.lower().startswith("charset=")
            ),
            None,
        )
        if media not in READABLE:
            raise UnreadableDocument(
                f"{url} serves {media}, which the library does not read; it reads "
                f"{', '.join(sorted(READABLE))}"
            )
        return media, charset
    from urllib.parse import urlsplit

    return media_type_of_name(urlsplit(url).path or url), None


def _text(content: bytes, charset: str | None, name: str) -> str:
    encoding = charset or "utf-8"
    try:
        return content.decode(encoding)
    except (UnicodeDecodeError, LookupError) as exc:
        raise UnreadableDocument(f"{name!r} is not {encoding} text: {exc}") from None


def read(content: bytes, media_type: str, name: str, charset: str | None = None) -> Document:
    """``content`` read as a document of ``media_type``; ``name`` (its file name or URL) names it
    in a refusal."""
    if media_type == HTML:
        document = _read_html(content, name, charset)
    elif media_type == MARKDOWN:
        document = _read_markdown(_text(content, charset, name))
    elif media_type == PDF:
        document = _read_pdf(content, name)
    elif media_type == WORD:
        document = _read_word(content, name)
    elif media_type == PLAIN:
        document = Document(title=None, sections=[_paragraphs(None, _text(content, charset, name))])
    else:
        raise UnreadableDocument(f"the library does not read {media_type} documents ({name!r})")
    document.sections = [s for s in document.sections if s.paragraphs]
    if not document.sections:
        raise UnreadableDocument(f"{name!r} holds no text")
    return document


_BLANK_LINES = re.compile(r"\n\s*\n")
_SPACE = re.compile(r"\s+")


def _clean(text: str) -> str:
    return _SPACE.sub(" ", text).strip()


def _paragraphs(heading: str | None, text: str) -> Section:
    return Section(
        heading=heading,
        page=None,
        paragraphs=[p for p in (_clean(b) for b in _BLANK_LINES.split(text)) if p],
    )


_MD_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")


def _read_markdown(text: str) -> Document:
    title: str | None = None
    sections: list[Section] = []
    heading: str | None = None
    lines: list[str] = []
    for line in text.splitlines():
        match = _MD_HEADING.match(line)
        if match is None:
            lines.append(line)
            continue
        sections.append(_paragraphs(heading, "\n".join(lines)))
        heading, lines = _clean(match.group(2)), []
        if title is None and len(match.group(1)) == 1:
            title = heading
    sections.append(_paragraphs(heading, "\n".join(lines)))
    return Document(title=title, sections=sections)


_HEADINGS = frozenset({"h1", "h2", "h3", "h4", "h5", "h6"})
_BLOCKS = frozenset({"p", "li", "pre", "blockquote", "td", "th", "dt", "dd", "figcaption"})
_DROPPED = frozenset(
    {"script", "style", "noscript", "template", "nav", "header", "footer", "aside", "form"}
)


def _read_html(content: bytes, name: str, charset: str | None) -> Document:
    import lxml.html
    from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]  # lxml ships no stubs for its C module

    try:
        root = (
            lxml.html.document_fromstring(_text(content, charset, name))
            if charset
            else lxml.html.document_fromstring(content)
        )
    except (etree.ParserError, ValueError) as exc:
        raise UnreadableDocument(f"{name!r} is not HTML that can be read: {exc}") from None
    for dropped in root.xpath("//*[" + " or ".join(f"self::{t}" for t in _DROPPED) + "]"):
        dropped.drop_tree()
    for comment in root.xpath("//comment()"):
        comment.drop_tree()
    titles = root.xpath("//title")
    title = _clean(titles[0].text_content()) if titles else None
    sections: list[Section] = [Section(heading=None, page=None)]
    body = root.find("body")
    for element in (body if body is not None else root).iter():
        tag = element.tag if isinstance(element.tag, str) else ""
        if tag in _HEADINGS:
            heading = _clean(element.text_content())
            sections.append(Section(heading=heading or None, page=None))
            if not title and tag == "h1":
                title = heading
        elif tag in _BLOCKS and not any(
            isinstance(a.tag, str) and a.tag in _BLOCKS for a in element.iterancestors()
        ):
            text = _clean(element.text_content())
            if text:
                sections[-1].paragraphs.append(text)
    return Document(title=title or None, sections=sections)


def _read_pdf(content: bytes, name: str) -> Document:
    """A PDF by its pages: each page a section, its paragraphs the page's blank-line runs."""
    import io

    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise UnreadableDocument(f"{name!r} is an encrypted PDF")
        pages = [page.extract_text() or "" for page in reader.pages]
        metadata = reader.metadata
    except PdfReadError as exc:
        raise UnreadableDocument(f"{name!r} is not a PDF that can be read: {exc}") from None
    sections = []
    for number, text in enumerate(pages, start=1):
        section = _paragraphs(None, text)
        section.page = number
        sections.append(section)
    title = _clean(metadata.title) if metadata is not None and metadata.title else None
    return Document(title=title or None, sections=sections)


def _read_word(content: bytes, name: str) -> Document:
    """A Word (.docx) document by its heading styles, its tables where they stand; a ``Title``
    paragraph titles it when its properties do not."""
    import io
    import zipfile

    import docx
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table

    try:
        word = docx.Document(io.BytesIO(content))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise UnreadableDocument(
            f"{name!r} is not a Word document that can be read: {exc}"
        ) from None
    title = _clean(word.core_properties.title or "") or None
    sections: list[Section] = [Section(heading=None, page=None)]
    for block in word.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                text = " | ".join(c for c in (_clean(cell.text) for cell in row.cells) if c)
                if text:
                    sections[-1].paragraphs.append(text)
            continue
        text = _clean(block.text)
        if not text:
            continue
        style = ""
        if block.style is not None:
            if block.style.name is None:
                raise UnreadableDocument(
                    f"{name!r} is not a Word document that can be read: "
                    f"a paragraph uses a style with no name ({block.style.style_id!r})"
                )
            style = block.style.name
        if style == "Title":
            title = title or text
        elif style.startswith("Heading"):
            sections.append(Section(heading=text, page=None))
        else:
            sections[-1].paragraphs.append(text)
    return Document(title=title, sections=sections)
