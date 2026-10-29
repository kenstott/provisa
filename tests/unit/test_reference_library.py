# Copyright (c) 2026 Kenneth Stott
# Canary: 909702f3-d57a-40e2-9dac-b083d69153bb
#
# This source code is licensed under the Business Source License 1.1
# found in the LICENSE file in the root directory of this source tree.
#
# NOTICE: Use of this software for training artificial intelligence or
# machine learning models is strictly prohibited without explicit written
# permission from the copyright holder.

"""REQ-2185: the reference library -- documents read into passages, embedded, held in the org's
control plane and searched for a role, each passage with what to cite; and the Admin page's API
that adds, keeps, rescopes, refreshes, replaces and removes them."""

# Requirements: REQ-2185

from __future__ import annotations

import base64
import re
import types
from typing import TYPE_CHECKING, Any, cast

import pytest
from sqlalchemy import insert, select

from provisa.api.admin import reference_library_router as router
from provisa.api.errors import ApiError
from provisa.core.schema_org import domains, reference_passages
from provisa.reference_library import extract, library, retrieve, store
from provisa.reference_library.passages import LIMIT, TARGET, cut

if TYPE_CHECKING:
    from fastapi import Request

SCHEMA = "org_default"

# -- reading a document -------------------------------------------------------------------------

POLICY = b"""<html><head><title>Client policy</title><script>var x=1;</script></head>
<body><nav><a href="/">Home</a></nav>
<h1>Clients</h1><p>A Client is a party that holds an account with the bank.</p>
<h2>Retention</h2><p>Client records are kept for seven years.</p>
<ul><li>Statements</li><li>Contracts <b>and</b> notices</li></ul>
<footer>Copyright</footer></body></html>"""

SPEC = b"""# Payments standard

Every payment carries an amount and a currency.

## Settlement

A payment settles on the second business day.
"""


def test_html_is_read_by_its_headings_without_its_chrome():
    document = extract.read(POLICY, extract.HTML, "policy.html")
    assert document.title == "Client policy"
    assert [(s.heading, s.paragraphs) for s in document.sections] == [
        ("Clients", ["A Client is a party that holds an account with the bank."]),
        (
            "Retention",
            ["Client records are kept for seven years.", "Statements", "Contracts and notices"],
        ),
    ]


def test_markdown_is_read_by_its_headings_and_titled_by_its_first():
    document = extract.read(SPEC, extract.MARKDOWN, "spec.md")
    assert document.title == "Payments standard"
    assert [s.heading for s in document.sections] == ["Payments standard", "Settlement"]


def test_plain_text_is_read_as_paragraphs():
    document = extract.read(b"one\nline\n\ntwo", extract.PLAIN, "notes.txt")
    assert document.title is None
    assert document.sections[0].paragraphs == ["one line", "two"]


@pytest.mark.parametrize(
    ("content", "media", "name", "says"),
    [
        (b"x", "application/zip", "a.zip", "does not read application/zip"),
        (b"\xff\xfe\x00", extract.PLAIN, "a.txt", "is not utf-8 text"),
        (b"<html><body><script>x</script></body></html>", extract.HTML, "a.html", "holds no text"),
    ],
)
def test_a_document_that_cannot_be_read_is_refused_by_name(content, media, name, says):
    with pytest.raises(extract.UnreadableDocument, match=says):
        extract.read(content, media, name)


def test_a_file_is_read_as_its_name_says_and_an_unknown_kind_is_refused():
    assert extract.media_type_of_name("Policy.HTML") == extract.HTML
    assert extract.media_type_of_name("spec.md") == extract.MARKDOWN
    assert extract.media_type_of_name("contract.PDF") == extract.PDF
    assert extract.media_type_of_name("standard.docx") == extract.WORD
    with pytest.raises(extract.UnreadableDocument, match="'deck.pptx' is not a document"):
        extract.media_type_of_name("deck.pptx")


def test_a_response_is_read_as_its_content_type_says():
    assert extract.media_type_of_response("text/html; charset=ISO-8859-1", "https://x/p") == (
        extract.HTML,
        "ISO-8859-1",
    )
    assert extract.media_type_of_response(None, "https://x/a/spec.md?v=1") == (
        extract.MARKDOWN,
        None,
    )
    with pytest.raises(extract.UnreadableDocument, match="serves image/png"):
        extract.media_type_of_response("image/png", "https://x/logo")


def _pdf(pages: list[str], title: str) -> bytes:
    """A PDF of ``pages``, each one line of Helvetica text, with ``title`` in its information."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append(f"<< /Title ({title}) >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info {len(objects)} 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return out


def test_a_pdf_is_read_by_its_pages():
    document = extract.read(
        _pdf(["A Client holds an account.", "Records are kept seven years."], "Client policy"),
        extract.PDF,
        "policy.pdf",
    )
    assert document.title == "Client policy"
    assert [(s.page, s.paragraphs) for s in document.sections] == [
        (1, ["A Client holds an account."]),
        (2, ["Records are kept seven years."]),
    ]
    assert [(p.page, p.section) for p in cut(document)] == [(1, None), (2, None)]


def test_a_word_document_is_read_by_its_headings_with_its_tables_in_place():
    import io

    import docx

    word = docx.Document()
    word.add_paragraph("Client policy", style="Title")
    word.add_heading("Clients", 1)
    word.add_paragraph("A Client holds an account.")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "Client"
    word.add_heading("Retention", 1)
    word.add_paragraph("Records are kept seven years.")
    data = io.BytesIO()
    word.save(data)
    document = extract.read(data.getvalue(), extract.WORD, "policy.docx")
    assert document.title == "Client policy"
    assert [(s.heading, s.paragraphs) for s in document.sections] == [
        ("Clients", ["A Client holds an account.", "Term | Client"]),
        ("Retention", ["Records are kept seven years."]),
    ]


@pytest.mark.parametrize(
    ("media", "name", "says"),
    [
        (extract.PDF, "x.pdf", "'x.pdf' is not a PDF that can be read"),
        (extract.WORD, "x.docx", "'x.docx' is not a Word document that can be read"),
    ],
)
def test_a_pdf_or_word_file_that_is_not_one_is_refused_by_name(media, name, says):
    with pytest.raises(extract.UnreadableDocument, match=says):
        extract.read(b"not a document", media, name)


def test_passages_keep_to_one_section_and_a_long_paragraph_is_cut_at_its_sentences():
    long = " ".join(f"Sentence {i} is about settlement terms." for i in range(120))
    document = extract.Document(
        title=None,
        sections=[
            extract.Section("A", None, ["short one", "short two"]),
            extract.Section("B", None, [long]),
        ],
    )
    passages = cut(document)
    assert passages[0].section == "A" and passages[0].text == "short one\n\nshort two"
    rest = passages[1:]
    assert {p.section for p in rest} == {"B"}
    assert all(len(p.text) <= max(LIMIT, TARGET) for p in rest)
    assert all(p.text.rstrip().endswith(".") for p in rest)
    assert [p.ordinal for p in passages] == list(range(len(passages)))


# -- the library in the control plane -----------------------------------------------------------

VOCABULARY = ("client", "account", "record", "payment", "currency", "settle", "contract")


def _vector(text: str) -> list[float]:
    words = re.findall(r"[a-z]+", text.lower())
    return [float(sum(w.startswith(v) for w in words)) + 0.01 for v in VOCABULARY]


async def _embed(texts: list[str]) -> list[list[float]]:
    return [_vector(t) for t in texts]


EMBEDDER = library.Embedder(model="test-embed", embed=_embed)


@pytest.fixture
def db(control_plane):
    with control_plane.engine().begin() as c:
        c.exec_driver_sql(f'SET LOCAL search_path TO "{SCHEMA}"')
        c.execute(insert(domains).values(id="sales", description=""))
        c.execute(insert(domains).values(id="finance", description=""))
    return control_plane.org()


def _file(name: str, data: bytes) -> library.Content:
    return library.file_content(name, data)


async def _add(conn, name: str, data: bytes, scope: list[str], status: str = store.KEPT) -> int:
    return await library.add(
        conn,
        _file(name, data),
        source_kind=store.FILE,
        domains=scope,
        status=status,
        embedder=EMBEDDER,
        added_by="alice",
    )


async def test_a_role_draws_on_the_documents_whose_domains_it_reads_and_every_whole_org_one(db):
    async with db.acquire() as conn:
        policy = await _add(conn, "policy.html", POLICY, ["sales"])
        await _add(conn, "spec.md", SPEC, ["finance"])
        whole = await _add(conn, "glossary.txt", b"A contract binds a client.", ["*"])

        sales = await retrieve.retrieve(conn, "What is a client?", frozenset({"sales"}), EMBEDDER)
        assert {c.document_id for c in sales} == {policy, whole}
        top = sales[0]
        assert (top.document_id, top.document_title, top.source, top.section) == (
            policy,
            "Client policy",
            "policy.html",
            "Clients",
        )
        assert top.quote == "A Client is a party that holds an account with the bank."

        finance = await retrieve.retrieve(conn, "client", frozenset({"finance"}), EMBEDDER)
        assert policy not in {c.document_id for c in finance}

        everyone = await retrieve.retrieve(conn, "payment currency", None, EMBEDDER, k=1)
        assert [c.document_title for c in everyone] == ["Payments standard"]


async def test_a_proposed_document_is_not_drawn_on_until_it_is_kept(db):
    async with db.acquire() as conn:
        proposed = await _add(conn, "policy.html", POLICY, ["*"], status=store.PROPOSED)
        assert await retrieve.retrieve(conn, "client", None, EMBEDDER) == []
        await library.keep(conn, proposed)
        assert {c.document_id for c in await retrieve.retrieve(conn, "client", None, EMBEDDER)} == {
            proposed
        }
        with pytest.raises(library.LibraryRefused, match="already kept"):
            await library.keep(conn, proposed)


async def test_a_document_already_in_the_library_is_not_added_again(db):
    async with db.acquire() as conn:
        first = await _add(conn, "policy.html", POLICY, ["sales"])
        with pytest.raises(library.AlreadyInLibrary) as refused:
            await _add(conn, "copy-of-policy.html", POLICY, ["finance"])
        assert refused.value.document_id == first
        assert len(await store.documents(conn)) == 1


async def test_a_new_version_replaces_the_content_and_keeps_place_status_and_domains(db):
    async with db.acquire() as conn:
        document_id = await _add(conn, "policy.html", POLICY, ["sales"])
        before = await store.document(conn, document_id)
        await library.new_version(
            conn, document_id, _file("policy-v2.md", b"# Policy\n\nA payment settles."), EMBEDDER
        )
        after = await store.document(conn, document_id)
        assert after is not None and before is not None
        assert (after.title, after.status, after.domains, after.added_at) == (
            before.title,
            before.status,
            ("sales",),
            before.added_at,
        )
        assert after.source == "policy-v2.md" and after.replaced_at is not None
        assert after.refreshed_at is None
        texts = [t for _, _, _, t in await store.passages(conn, document_id)]
        assert texts == ["A payment settles."]
        with pytest.raises(library.LibraryRefused, match="uploaded as a file"):
            await library.refresh(conn, document_id, EMBEDDER, None)


async def test_refresh_fetches_a_url_document_again(db, monkeypatch):
    served = {"body": SPEC}

    async def fetch(url, org_id):
        return library.Content(served["body"], extract.MARKDOWN, None, url)

    monkeypatch.setattr(library, "fetch", fetch)
    url = "https://example.com/payments.md"
    async with db.acquire() as conn:
        document_id = await library.add(
            conn,
            await library.fetch(url, None),
            source_kind=store.URL,
            domains=["finance"],
            status=store.KEPT,
            embedder=EMBEDDER,
            added_by=None,
        )
        with pytest.raises(library.AlreadyInLibrary):
            await library.add(
                conn,
                await library.fetch(url, None),
                source_kind=store.URL,
                domains=["finance"],
                status=store.KEPT,
                embedder=EMBEDDER,
                added_by=None,
            )
        served["body"] = b"# Payments standard\n\nA payment settles on the same day."
        await library.refresh(conn, document_id, EMBEDDER, None)
        held = await store.document(conn, document_id)
        assert held is not None and held.refreshed_at is not None and held.source == url
        assert [t for *_, t in await store.passages(conn, document_id)] == [
            "A payment settles on the same day."
        ]
        with pytest.raises(library.LibraryRefused, match="added from its URL"):
            await library.new_version(conn, document_id, _file("x.txt", b"other"), EMBEDDER)


async def test_removing_a_document_takes_its_passages_with_it(db):
    async with db.acquire() as conn:
        document_id = await _add(conn, "policy.html", POLICY, ["*"])
        await library.remove(conn, document_id)
        assert await retrieve.retrieve(conn, "client", None, EMBEDDER) == []
        held = (await conn.execute_core(select(reference_passages.c.id))).fetchall()
        assert held == []
        with pytest.raises(library.DocumentNotFound):
            await library.remove(conn, document_id)


async def test_a_document_embedded_by_another_model_is_named_rather_than_compared(db):
    async with db.acquire() as conn:
        await _add(conn, "policy.html", POLICY, ["*"])
        other = library.Embedder(model="other-embed", embed=_embed)
        with pytest.raises(retrieve.LibraryModelChanged, match="'Client policy'.*'other-embed'"):
            await retrieve.retrieve(conn, "client", None, other)


@pytest.mark.parametrize(
    ("raw", "says"),
    [
        ([], "list of domain ids"),
        ("sales", "list of domain ids"),
        (["*", "sales"], "cannot be combined"),
        (["marketing"], "no domain 'marketing'"),
    ],
)
def test_a_documents_domains_are_the_orgs_or_the_whole_org(raw, says):
    with pytest.raises(library.LibraryRefused, match=says):
        library.check_domains(raw, {"sales", "finance"})
    assert library.check_domains(["finance", "sales", "sales"], {"sales", "finance"}) == [
        "finance",
        "sales",
    ]
    assert library.check_domains(["*"], set()) == ["*"]


def test_a_url_is_an_http_address():
    assert library.check_url("https://example.com/a") == "https://example.com/a"
    for bad in ("ftp://example.com/a", "example.com/a", "https://", 3):
        with pytest.raises(library.LibraryRefused, match="http"):
            library.check_url(bad)


# -- the Admin page's API -----------------------------------------------------------------------


def _request() -> Request:
    fake = types.SimpleNamespace(
        state=types.SimpleNamespace(
            identity=types.SimpleNamespace(user_id="alice", roles=[], acting_roles=[]),
            active_org_id="default",
        )
    )
    return cast("Request", fake)


@pytest.fixture
def api(db, monkeypatch):
    monkeypatch.setattr(router, "_pool", lambda: db)
    monkeypatch.setattr(router, "_guard", lambda request: "default")

    async def embedder():
        return EMBEDDER

    monkeypatch.setattr(router, "_embedder", embedder)
    return db


def _upload(name: str, data: bytes, **extra: Any) -> dict:
    return {"file_name": name, "content_b64": base64.b64encode(data).decode(), **extra}


async def test_the_admin_api_adds_lists_rescopes_and_removes_a_document(api):
    added = await router.add_document(_upload("policy.html", POLICY, domains=["sales"]), _request())
    assert (added["title"], added["source_kind"], added["status"], added["domains"]) == (
        "Client policy",
        "file",
        "kept",
        ["sales"],
    )
    assert added["passages"] == 2 and added["added_by"] == "alice"
    listed = await router.list_documents(_request())
    assert [d["id"] for d in listed["documents"]] == [added["id"]]

    changed = await router.change_document(
        added["id"], {"domains": ["*"], "title": "Clients"}, _request()
    )
    assert (changed["domains"], changed["title"]) == (["*"], "Clients")

    detail = await router.get_document(added["id"], _request())
    assert [p["section"] for p in detail["passage_list"]] == ["Clients", "Retention"]

    assert await router.remove_document(added["id"], _request()) == {"removed": added["id"]}
    with pytest.raises(ApiError) as gone:
        await router.get_document(added["id"], _request())
    assert gone.value.status_code == 404 and gone.value.code == "reference_library.not_found"


async def test_the_admin_api_refuses_by_name(api):
    async def refused(body: dict) -> ApiError:
        with pytest.raises(ApiError) as caught:
            await router.add_document(body, _request())
        return caught.value

    unknown = await refused(_upload("policy.html", POLICY, domains=["marketing"]))
    assert unknown.status_code == 400 and "marketing" in str(unknown.detail)
    both = await refused({**_upload("a.txt", b"x"), "url": "https://x/a", "domains": ["*"]})
    assert "not both" in str(both.detail)
    unreadable = await refused(_upload("deck.pptx", b"x", domains=["*"]))
    assert "'deck.pptx' is not a document" in str(unreadable.detail)
    added = await router.add_document(_upload("policy.html", POLICY, domains=["*"]), _request())
    again = await refused(_upload("policy.html", POLICY, domains=["*"]))
    assert again.status_code == 409 and again.params == {"document_id": added["id"]}
    with pytest.raises(ApiError, match="discarding it removes it"):
        await router.change_document(added["id"], {"status": "proposed"}, _request())


async def test_the_admin_api_uploads_a_new_version(api):
    added = await router.add_document(_upload("spec.md", SPEC, domains=["finance"]), _request())
    replaced = await router.upload_version(
        added["id"], _upload("spec-v2.txt", b"Payments settle daily."), _request()
    )
    assert replaced["source"] == "spec-v2.txt" and replaced["replaced_at"] is not None
    assert replaced["passages"] == 1


# -- what an assistant reaches ------------------------------------------------------------------


async def test_the_mcp_tool_searches_for_the_callers_role(db, monkeypatch):
    from provisa.api.mcp import tools

    async with db.acquire() as conn:
        await _add(conn, "policy.html", POLICY, ["sales"])
        await _add(conn, "spec.md", SPEC, ["finance"])

    async def embedder(state):
        return EMBEDDER

    monkeypatch.setattr(retrieve, "org_embedder", embedder)
    state = types.SimpleNamespace(
        model_db=db,
        roles={"analyst": {"id": "analyst", "domain_access": ["sales"], "capabilities": []}},
    )
    monkeypatch.setattr(tools, "require_role", lambda role, state: role)
    found = await tools.search_reference_library(cast(Any, state), "analyst", "client account")
    assert {f["document_title"] for f in found} == {"Client policy"}
    assert set(found[0]) == {
        "document_id",
        "document_title",
        "source",
        "section",
        "page",
        "quote",
        "score",
    }
